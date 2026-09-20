from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from typing import Optional
from ..database import get_database
from ..utils.auth import require_admin
from bson import ObjectId
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from io import BytesIO

router = APIRouter(prefix="/api/reports", tags=["Reports"])

def create_excel_response(workbook, filename):
    """Helper to create Excel file response"""
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    
    headers = {
        'Content-Disposition': f'attachment; filename="{filename}"',
        'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    }
    
    return StreamingResponse(output, headers=headers, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@router.get("/loans")
async def generate_loans_report(
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    current_user: dict = Depends(require_admin),
    db = Depends(get_database)
):
    """Generate loans report with Excel export"""
    
    query = {}
    
    if start_date:
        query["createdAt"] = {"$gte": datetime.fromisoformat(start_date)}
    if end_date:
        if "createdAt" in query:
            query["createdAt"]["$lte"] = datetime.fromisoformat(end_date)
        else:
            query["createdAt"] = {"$lte": datetime.fromisoformat(end_date)}
    if status:
        query["status"] = status
    
    loans = await db.loans.find(query).sort("createdAt", -1).to_list(1000)
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Loans Report"
    
    # Headers
    headers = ["Borrower", "Loan Type", "Amount", "Interest Rate", "Total Repayment", "Remaining Balance", "Progress", "Status", "Applied Date", "Approval Date"]
    ws.append(headers)
    
    # Style headers
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    # Add data
    for loan in loans:
        user = await db.users.find_one({"_id": loan.get("userId")})
        ws.append([
            user.get("name", "Unknown") if user else "Unknown",
            loan.get("loanType", ""),
            loan.get("amount", 0),
            loan.get("interestRate", 0),
            loan.get("totalRepaymentAmount", 0),
            loan.get("remainingBalance", 0),
            f"{loan.get('paymentProgress', 0):.1f}%",
            loan.get("status", ""),
            loan.get("createdAt").strftime("%Y-%m-%d") if isinstance(loan.get("createdAt"), datetime) else "",
            loan.get("approvalDate").strftime("%Y-%m-%d") if isinstance(loan.get("approvalDate"), datetime) else ""
        ])
    
    # Auto-size columns
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    return create_excel_response(wb, f"loans_report_{datetime.now().strftime('%Y%m%d')}.xlsx")

@router.get("/payments")
async def generate_payments_report(
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    current_user: dict = Depends(require_admin),
    db = Depends(get_database)
):
    """Generate payments report"""
    
    query = {}
    
    if start_date:
        query["paymentDate"] = {"$gte": datetime.fromisoformat(start_date)}
    if end_date:
        if "paymentDate" in query:
            query["paymentDate"]["$lte"] = datetime.fromisoformat(end_date)
        else:
            query["paymentDate"] = {"$lte": datetime.fromisoformat(end_date)}
    
    payments = await db.loanhistories.find(query).sort("paymentDate", -1).to_list(5000)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Payments Report"
    
    headers = ["Borrower", "Amount Paid", "Payment Date", "Remaining Balance", "Total Repayment", "Notes"]
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    for payment in payments:
        user = await db.users.find_one({"_id": payment.get("userId")})
        ws.append([
            user.get("name", "Unknown") if user else "Unknown",
            payment.get("amountPaid", 0),
            payment.get("paymentDate").strftime("%Y-%m-%d") if isinstance(payment.get("paymentDate"), datetime) else "",
            payment.get("remainingBalance", 0),
            payment.get("totalRepaymentAmount", 0),
            payment.get("notes", "")
        ])
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    return create_excel_response(wb, f"payments_report_{datetime.now().strftime('%Y%m%d')}.xlsx")

@router.get("/savings")
async def generate_savings_report(
    current_user: dict = Depends(require_admin),
    db = Depends(get_database)
):
    """Generate savings report"""
    
    users_list = await db.users.find({"active": True}).sort("name", 1).to_list(1000)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Savings Report"
    
    headers = ["Member Name", "Email", "Phone", "Balance", "Last Updated"]
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="9966FF", end_color="9966FF", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    for user in users_list:
        account = await db.accountbalances.find_one({"userId": user["_id"]})
        ws.append([
            user.get("name", ""),
            user.get("email", ""),
            str(user.get("phoneNumber", "")),
            account.get("balanceAmount", 0) if account else 0,
            account.get("updatedAt").strftime("%Y-%m-%d") if account and isinstance(account.get("updatedAt"), datetime) else ""
        ])
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    return create_excel_response(wb, f"savings_report_{datetime.now().strftime('%Y%m%d')}.xlsx")

@router.get("/members")
async def generate_members_report(
    current_user: dict = Depends(require_admin),
    db = Depends(get_database)
):
    """Generate members report"""
    
    users_list = await db.users.find().sort("name", 1).to_list(1000)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Members Report"
    
    headers = ["Name", "Email", "Phone", "Address", "State", "LGA", "Role", "Status", "Next of Kin", "NOK Phone", "Joined Date"]
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    for user in users_list:
        ws.append([
            user.get("name", ""),
            user.get("email", ""),
            str(user.get("phoneNumber", "")),
            user.get("address", ""),
            user.get("homeState", ""),
            user.get("lga", ""),
            user.get("role", ""),
            "Active" if user.get("active") else "Inactive",
            user.get("nextOfKinName", ""),
            str(user.get("nextOfKinNumber", "")),
            user.get("createdAt").strftime("%Y-%m-%d") if isinstance(user.get("createdAt"), datetime) else ""
        ])
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    return create_excel_response(wb, f"members_report_{datetime.now().strftime('%Y%m%d')}.xlsx")

@router.get("/dividends")
async def generate_dividends_report(
    current_user: dict = Depends(require_admin),
    db = Depends(get_database)
):
    """Generate dividends report"""
    
    dividends = await db.dividends.find().sort("createdAt", -1).to_list(1000)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Dividends Report"
    
    headers = ["Title", "Total Amount", "Distribution Method", "Recipients", "Status", "Created Date", "Distributed Date"]
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    for dividend in dividends:
        ws.append([
            dividend.get("title", ""),
            dividend.get("totalAmount", 0),
            dividend.get("distributionMethod", ""),
            dividend.get("totalRecipients", 0),
            dividend.get("status", ""),
            dividend.get("createdAt").strftime("%Y-%m-%d") if isinstance(dividend.get("createdAt"), datetime) else "",
            dividend.get("distributedAt").strftime("%Y-%m-%d") if isinstance(dividend.get("distributedAt"), datetime) else ""
        ])
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    return create_excel_response(wb, f"dividends_report_{datetime.now().strftime('%Y%m%d')}.xlsx")

@router.get("/attendance")
async def generate_attendance_report(
    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),
    current_user: dict = Depends(require_admin),
    db = Depends(get_database)
):
    """Generate attendance report"""
    
    query = {}
    
    if start_date:
        query["meetingDate"] = {"$gte": datetime.fromisoformat(start_date)}
    if end_date:
        if "meetingDate" in query:
            query["meetingDate"]["$lte"] = datetime.fromisoformat(end_date)
        else:
            query["meetingDate"] = {"$lte": datetime.fromisoformat(end_date)}
    
    records = await db.attendance_records.find(query).sort("meetingDate", -1).to_list(1000)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    
    headers = ["Meeting Title", "Date", "Location", "Total Members", "Present", "Online", "Absent", "Total Points", "Attendance Rate"]
    ws.append(headers)
    
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    
    for record in records:
        ws.append([
            record.get("title", ""),
            record.get("meetingDate").strftime("%Y-%m-%d %H:%M") if isinstance(record.get("meetingDate"), datetime) else "",
            record.get("location", ""),
            record.get("totalMembers", 0),
            record.get("presentCount", 0),
            record.get("onlineCount", 0),
            record.get("absentCount", 0),
            f"{record.get('totalPoints', 0)}/{record.get('maxPoints', 0)}",
            f"{record.get('attendanceRate', 0):.1f}%"
        ])
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    return create_excel_response(wb, f"attendance_report_{datetime.now().strftime('%Y%m%d')}.xlsx")

@router.get("/member-attendance/{user_id}")
async def generate_member_attendance_report(
    user_id: str,
    current_user: dict = Depends(require_admin),
    db = Depends(get_database)
):
    """Generate individual member attendance report"""
    
    try:
        user_obj_id = ObjectId(user_id)
    except:
        raise HTTPException(status_code=400, detail="Invalid user ID")
    
    user = await db.users.find_one({"_id": user_obj_id})
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    marks = await db.attendance_marks.find({"userId": user_obj_id}).sort("markedAt", -1).to_list(1000)
    
    wb = Workbook()
    ws = wb.active
    ws.title = f"{user.get('name', 'Member')} Attendance"
    
    # Add member info
    ws.append(["Member Attendance Report"])
    ws.append([f"Name: {user.get('name', '')}", f"Email: {user.get('email', '')}"])
    ws.append([])
    
    headers = ["Meeting Title", "Meeting Date", "Location", "Status", "Points", "Marked Date"]
    ws.append(headers)
    
    for cell in ws[4]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    
    for mark in marks:
        record = await db.attendance_records.find_one({"_id": mark.get("attendanceId")})
        if record:
            points_map = {"present": 2, "online": 1, "absent": 0}
            ws.append([
                record.get("title", ""),
                record.get("meetingDate").strftime("%Y-%m-%d") if isinstance(record.get("meetingDate"), datetime) else "",
                record.get("location", ""),
                mark.get("status", ""),
                points_map.get(mark.get("status", "absent"), 0),
                mark.get("markedAt").strftime("%Y-%m-%d") if isinstance(mark.get("markedAt"), datetime) else ""
            ])
    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
    
    return create_excel_response(wb, f"{user.get('name', 'member')}_attendance_{datetime.now().strftime('%Y%m%d')}.xlsx")